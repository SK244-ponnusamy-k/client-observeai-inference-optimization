"""
Shared QID dataset parser.

This is the SINGLE source of truth for turning an AutoQA dataset (xlsx / csv /
jsonl) into per-row dicts, so both the quality evaluator (inference/quality-eval.py)
and the benchmark harness (inference/load-test.py) parse identically — same prompt
assembly, same QID grouping, same ORG/XL sheet handling.

It was factored out of quality-eval.py (which has a hyphen in its name and so is
not importable as a module). quality-eval.py now imports from here; the benchmark's
QID-cost mode imports the same functions so a per-QID cost run segments the dataset
exactly the way the quality run scored it.

Row shape returned by load_dataset():
    {
      "data_id":       stable id (or row index),
      "prompt":        the fully assembled model input (transcript+question or
                       the prebuilt input_prompt column),
      "gold":          ground-truth label normalised to "Yes"/"No",
      "question":      grouping key — the QID when present, else a 60-char stub,
      "question_id":   the stable QID on its own ("" if absent),
      "question_full": the full rubric question text,
      "prompt_chars":  len(prompt) as a string (transcript size WITHOUT persisting it),
    }
"""

from __future__ import annotations

import csv
import json
import logging
import re
import sys
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)
csv.field_size_limit(10_000_000)  # transcripts are large


# ---------------------------------------------------------------------------
# Dataset loading (xlsx / csv / jsonl)
# ---------------------------------------------------------------------------
def load_dataset(path: str, cfg: dict[str, Any], sheet: str = "") -> list[dict[str, str]]:
    dcol = cfg["dataset"]
    text_f = dcol["text_field"]
    label_f = dcol["label_field"]
    q_f = dcol.get("question_field", "")
    # Column holding the stable question id (QID). When present it is the
    # PREFERRED grouping key for per-question metrics, because it is stable across
    # the two dataset sheets (Test-ORG / Test-XL) whereas the question text can be
    # reworded. Falls back to the question text when absent (older datasets).
    qid_f = dcol.get("question_id_field", "question_id")
    # Column holding the raw call transcript (context). When set, the prompt sent
    # to the model is ASSEMBLED from this transcript plus the QID question text
    # (the agent instruction), instead of using a prebuilt input_prompt column.
    transcript_f = dcol.get("transcript_field", "")
    # Sheet-name overrides (xlsx). ``sheet`` (arg) selects the transcript sheet to
    # score; ``qid_sheet`` is the lookup table of question_id -> question text.
    transcript_sheet = sheet or dcol.get("sheet", "")
    qid_sheet = dcol.get("qid_sheet", "QIDs")
    qid_id_col = dcol.get("qid_id_column", "Question Id")
    qid_text_col = dcol.get("qid_text_column", "Question Text")
    # Template used to build the prompt from transcript + question. {question} is
    # the QID rubric text (instruction); {transcript} is the call context.
    prompt_template = dcol.get("prompt_template") or (
        "Evaluate the human agent against the given question and transcript.\n"
        "=== QUESTION ===\n{question}\n=== TRANSCRIPT ===\n{transcript}"
    )
    rows: list[dict[str, str]] = []
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"Dataset not found: {path}")

    is_xlsx = p.suffix.lower() in (".xlsx", ".xlsm")

    # QID -> question-text lookup, populated from the QIDs sheet for xlsx sources.
    # Empty for csv/jsonl, where the question travels with each row.
    qid_map: dict[str, str] = _load_qid_map(p, qid_sheet, qid_id_col, qid_text_col) if is_xlsx else {}

    def _row(d: dict[str, Any], i: int) -> dict[str, str]:
        lower_d = {str(k).strip().lower(): v for k, v in d.items() if k is not None}
        gold_val = (
            d.get(label_f) or lower_d.get(label_f.lower()) or
            lower_d.get("answer") or lower_d.get("gold") or lower_d.get("ground_truth") or lower_d.get("label") or ""
        )
        q_val = d.get(q_f) or lower_d.get(q_f.lower()) or ""
        qid_val = d.get(qid_f) or lower_d.get(qid_f.lower()) or ""
        transcript_val = (d.get(transcript_f) or lower_d.get(transcript_f.lower()) or "") if transcript_f else ""

        gold_str = str(gold_val).strip()
        if gold_str.lower() == "yes":
            gold_str = "Yes"
        elif gold_str.lower() == "no":
            gold_str = "No"

        q_str = str(q_val).strip()
        qid_str = str(qid_val).strip()
        transcript_str = str(transcript_val).strip()

        # Resolve the rubric question. Prefer the QID lookup (authoritative full
        # text incl. sub-criteria); fall back to the row's own question column.
        if qid_str and qid_map.get(qid_str):
            q_str = qid_map[qid_str].strip()

        # Assemble the model prompt.
        #   1. transcript + QID question  -> build from the template (new xlsx)
        #   2. otherwise                   -> the prebuilt input_prompt column
        if transcript_str and q_str:
            prompt_str = prompt_template.format(question=q_str, transcript=transcript_str).strip()
        else:
            prompt_val = (
                d.get(text_f) or lower_d.get(text_f.lower()) or
                lower_d.get("input_prompt") or lower_d.get("question") or lower_d.get("prompt") or ""
            )
            prompt_str = str(prompt_val).strip()

        # If we still have no separate question, recover it from the prompt text
        # so the per-question grouping and report label are populated.
        if (not q_str or q_str == prompt_str) and "Question:" in prompt_str:
            q_match = re.search(r"Question:\s*(.*?)(?:\n|Sub-criteria:|$)", prompt_str, re.IGNORECASE | re.DOTALL)
            if q_match:
                q_str = q_match.group(1).strip()

        # Grouping key for per-question metrics. Prefer the stable QID (e.g.
        # "QID_1"); it is identical across Test-ORG and Test-XL, so the two
        # sheets' per-question breakdowns line up. Fall back to a compact slice
        # of the question text when no QID column is present.
        group_key = qid_str or (q_str[:60] if q_str else "all")

        return {
            "data_id": str(d.get("data_id") or lower_d.get("data_id") or i),
            "prompt": prompt_str,
            "gold": gold_str,
            # Key used to GROUP per-question metrics. It is a label, not the model's
            # input: the QID when available, else a 60-char question stub.
            "question": group_key,
            # Stable question id on its own, carried through to the sample dump.
            "question_id": qid_str or "",
            # Full rubric question, carried through to the sample dump so the
            # report shows the criterion in full rather than a 60-char stub.
            # This is the QA rubric, not customer data.
            "question_full": q_str or "all",
            # Size of the real prompt sent to the model (transcript + question +
            # sub-criteria). Recorded so a reader can see how large the actual
            # input was WITHOUT persisting the transcript itself.
            "prompt_chars": str(len(prompt_str)),
        }

    if is_xlsx:
        rows = [_row(d, i) for i, d in enumerate(_read_xlsx_sheet(p, transcript_sheet))]
    else:
        # Auto-detect JSONL content by checking first non-empty characters
        first_chars = ""
        with p.open(encoding="utf-8") as f:
            for line in f:
                s = line.strip()
                if s:
                    first_chars = s[:10]
                    break

        is_json = p.suffix.lower() == ".jsonl" or first_chars.startswith("{") or first_chars.startswith("[")

        if is_json:
            with p.open(encoding="utf-8") as f:
                for i, line in enumerate(f):
                    line = line.strip()
                    if line:
                        try:
                            rows.append(_row(json.loads(line), i))
                        except Exception as exc:  # noqa: BLE001
                            logger.warning("Failed to parse JSONL line %d: %s", i + 1, exc)
        else:  # csv / tsv
            with p.open(newline="", encoding="utf-8") as f:
                sample = f.read(4096)
                f.seek(0)
                first_line = sample.splitlines()[0] if sample else ""
                delim = "\t" if "\t" in first_line else ","
                rows = [_row(d, i) for i, d in enumerate(csv.DictReader(f, delimiter=delim))]

    n = int(cfg.get("run", {}).get("max_samples", 0) or 0)
    if n > 0:
        rows = rows[:n]
    return rows


# ---------------------------------------------------------------------------
# xlsx readers (openpyxl). Auto-installed the same way report_generator does, so
# the in-cluster pod needs no extra baked-in dependency.
# ---------------------------------------------------------------------------
def _import_openpyxl():  # type: ignore[no-untyped-def]
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        import subprocess  # noqa: PLC0415

        logger.info("openpyxl not present — installing it to read the .xlsx dataset.")
        subprocess.run([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl"], check=False)
        import openpyxl  # noqa: PLC0415
    return openpyxl


def _read_xlsx_sheet(path: Path, sheet: str) -> list[dict[str, Any]]:
    """Read one worksheet into a list of {header: value} dicts.

    The first row is the header. ``sheet`` selects the worksheet by name; when
    empty the workbook's first sheet is used. read_only + values_only keeps memory
    flat for the large XL transcripts.

    data_only is FALSE on purpose. These transcripts frequently start with
    "===== CALL n =====", which openpyxl classifies as a formula cell; with
    data_only=True such a cell returns None (no cached result) and the transcript
    would be dropped. Reading raw values returns the literal text instead. The
    dataset carries no genuine spreadsheet formulas (Word/Token counts are
    literals), so nothing computed is lost.
    """
    openpyxl = _import_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet:
            if sheet not in wb.sheetnames:
                raise KeyError(
                    f"Sheet '{sheet}' not found in {path.name}. "
                    f"Available: {', '.join(wb.sheetnames)}"
                )
            ws = wb[sheet]
        else:
            ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
        except StopIteration:
            return []
        out: list[dict[str, Any]] = []
        for values in rows_iter:
            if values is None or all(v is None for v in values):
                continue
            record = {header[i]: values[i] for i in range(min(len(header), len(values)))}
            out.append(record)
        return out
    finally:
        wb.close()


def _load_qid_map(path: Path, sheet: str, id_col: str, text_col: str) -> dict[str, str]:
    """Build a ``question_id -> question_text`` map from the QIDs sheet.

    Returns an empty map (and logs a warning) when the sheet is absent, so a
    workbook without a QIDs sheet degrades to using each row's own question
    column rather than failing the whole run.
    """
    openpyxl = _import_openpyxl()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet not in wb.sheetnames:
            logger.warning("QID sheet '%s' not found in %s; using per-row questions.", sheet, path.name)
            return {}
    finally:
        wb.close()

    mapping: dict[str, str] = {}
    for record in _read_xlsx_sheet(path, sheet):
        lower = {str(k).strip().lower(): v for k, v in record.items() if k}
        qid = record.get(id_col) or lower.get(id_col.lower()) or lower.get("question_id") or ""
        text = record.get(text_col) or lower.get(text_col.lower()) or lower.get("question_text") or ""
        qid_s = str(qid).strip()
        text_s = str(text).strip()
        if qid_s and text_s:
            mapping[qid_s] = text_s
    logger.info("Loaded %d QID -> question mappings from sheet '%s'.", len(mapping), sheet)
    return mapping
