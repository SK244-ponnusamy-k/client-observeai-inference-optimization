#!/usr/bin/env python3
"""
scripts/compare_models.py

Parses all JSONL benchmark results from results/ directory (recursively)
or syncs from S3, then generates a multi-model comparison table and Excel/CSV report.

Usage:
  python scripts/compare_models.py
  python report-gen/report_generator.py --date today
  python report-gen/report_generator.py --date 2026-09-15
  python report-gen/report_generator.py --from-date 2026-09-18 --to-date today
  python report-gen/report_generator.py --s3-bucket shellkode-ai-results
  python scripts/compare_models.py --export-excel results/model_comparison.xlsx
"""

import argparse
import csv
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


def load_results_local(results_dir: Path) -> list[dict[str, Any]]:
    rows = []
    if not results_dir.exists():
        return rows

    for f in results_dir.rglob("*.jsonl"):
        if (
            "quality" in f.parts
            or f.name.startswith("comparison")
            or f.name.startswith("summary")
            # Per-row quality samples carry no perf metrics; without this they
            # would pass the "no f1/accuracy key" test below and pollute the
            # performance sheet.
            or f.name.startswith("samples-")
        ):
            continue
        with f.open("r", encoding="utf-8") as file:
            for line in file:
                if line.strip():
                    try:
                        data = json.loads(line)
                        if "f1" not in data and "f1_score" not in data and "accuracy" not in data:
                            # Tag the row with its source path so date filtering can
                            # fall back to the S3 run-folder timestamp (results/<ts>/...)
                            # when the row's own started_at is on a different day.
                            data["_source_path"] = str(f)
                            rows.append(data)
                    except json.JSONDecodeError:
                        pass
    return rows


def sync_results_from_s3(bucket: str, local_dir: Path) -> None:
    """Sync both load benchmark results and quality eval results from S3."""
    import subprocess

    cmd_perf = f"aws s3 sync s3://{bucket}/results/ {local_dir}/perf/"
    cmd_qual = f"aws s3 sync s3://{bucket}/quality/ {local_dir}/quality/"
    print(f"Syncing benchmark & quality results from S3 (s3://{bucket}) ...")
    try:
        subprocess.run(cmd_perf, shell=True, check=True)
        subprocess.run(cmd_qual, shell=True, check=True)
        print("S3 sync complete.")
    except Exception as e:
        print(f"Warning: S3 sync failed or AWS CLI not found: {e}")


def load_quality_results(results_dir: Path) -> list[dict[str, Any]]:
    rows = []
    if not results_dir.exists():
        return rows

    for f in results_dir.rglob("*.jsonl"):
        if f.name.startswith("samples-"):
            continue  # per-row samples are loaded separately by load_quality_samples
        with f.open("r", encoding="utf-8") as file:
            for line in file:
                if line.strip():
                    try:
                        data = json.loads(line)
                        if "f1" in data or "f1_score" in data or "accuracy" in data:
                            data["_source_path"] = str(f)
                            model = data.get("served_model") or data.get("model") or data.get("hf_id") or "gpt-oss-20b"
                            data["_clean_model"] = model
                            data["_clean_hw"] = data.get("hardware") or data.get("instance_type") or "-"
                            data["_clean_quant"] = data.get("quantization") or "-"
                            data["_clean_rows"] = int(_v(data, "n_total", "total_rows", "n_scored"))
                            data["_clean_acc"] = _v(data, "accuracy") * (100.0 if _v(data, "accuracy") <= 1.0 else 1.0)
                            data["_clean_f1"] = _v(data, "f1", "f1_score", "macro_f1")
                            rows.append(data)
                    except json.JSONDecodeError:
                        pass
    return rows


def _run_label_from_path(path: Path) -> str:
    """'quality/20260922-104321/g5/samples-x.jsonl' -> '20260922-104321/g5'.

    Identifies the run by its timestamp folder plus any tag folder beneath it, so
    several runs of the same model on the same hardware stay distinguishable.
    """
    import re

    segs = str(path).replace("\\", "/").split("/")[:-1]  # drop the filename
    for i, seg in enumerate(segs):
        if re.match(r"^\d{8}-\d{6}$", seg.strip()):
            return "/".join(segs[i:])
    return segs[-1] if segs else "-"


def _model_from_samples_filename(name: str) -> str:
    """'samples-autoqa_v1-Qwen3.5-4B.jsonl' -> 'Qwen3.5-4B'.

    Only used for sample files written before the served_model field was added to
    each row; current files carry the model inline.
    """
    stem = name[len("samples-"):] if name.startswith("samples-") else name
    stem = stem.rsplit(".jsonl", 1)[0]
    parts = stem.split("-", 1)
    return parts[1] if len(parts) == 2 and parts[1] else stem or "-"


def load_quality_samples(results_dir: Path) -> list[dict[str, Any]]:
    """Read per-row quality samples written by quality-eval.py --dump-samples.

    Pure reader: every value shown in the report comes from the sample row as
    written by the evaluator. Nothing is recomputed here. Files are only present
    for runs that opted in, so an empty list simply means no run dumped samples.
    """
    rows: list[dict[str, Any]] = []
    if not results_dir.exists():
        return rows

    for f in sorted(results_dir.rglob("samples-*.jsonl")):
        fallback_model = _model_from_samples_filename(f.name)
        run_label = _run_label_from_path(f)
        with f.open("r", encoding="utf-8") as file:
            for line in file:
                if not line.strip():
                    continue
                try:
                    data = json.loads(line)
                except json.JSONDecodeError:
                    continue
                data["_source_path"] = str(f)
                data["_run"] = run_label
                data["_clean_model"] = data.get("served_model") or fallback_model
                data["_clean_hw"] = data.get("hardware") or "-"
                data["_clean_quant"] = data.get("quantization") or "-"
                rows.append(data)
    return rows


def _resolve_date_value(date_arg: str | None, option_name: str) -> str | None:
    """Resolve ``today`` or validate an ISO YYYY-MM-DD CLI date value."""
    if not date_arg:
        return None
    if date_arg.lower() == "today":
        return datetime.now(timezone.utc).strftime("%Y-%m-%d")
    try:
        datetime.strptime(date_arg, "%Y-%m-%d")
    except ValueError:
        raise SystemExit(
            f"Invalid {option_name} value: '{date_arg}'. Use 'today' or YYYY-MM-DD "
            "(e.g. 2026-09-18)."
        )
    return date_arg


def resolve_date_filter(date_arg: str | None) -> str | None:
    """Backward-compatible resolver for the exact ``--date`` filter."""
    return _resolve_date_value(date_arg, "--date")


def resolve_date_range(
    date_arg: str | None,
    from_date_arg: str | None,
    to_date_arg: str | None,
) -> tuple[str | None, str | None]:
    """Resolve exact-date or inclusive range arguments.

    ``--date`` is mutually exclusive with ``--from-date``/``--to-date``.
    When only ``--from-date`` is supplied, the upper bound defaults to today
    (UTC), making ``--from-date YYYY-MM-DD`` the convenient "through now" form.
    ``--to-date`` alone means all available results up to that date.
    """
    if date_arg and (from_date_arg or to_date_arg):
        raise SystemExit("Use either --date OR --from-date/--to-date, not both.")

    if date_arg:
        exact = resolve_date_filter(date_arg)
        return exact, exact

    from_date = _resolve_date_value(from_date_arg, "--from-date")
    to_date = _resolve_date_value(to_date_arg, "--to-date")
    if from_date and not to_date:
        to_date = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    if from_date and to_date and from_date > to_date:
        raise SystemExit(
            f"Invalid date range: --from-date {from_date} is after --to-date {to_date}."
        )
    return from_date, to_date


def _folder_date_from_path(path: str) -> str | None:
    """Extract the run-folder date from an S3-synced path as YYYY-MM-DD.

    Results are stored under results/<YYYYMMDD-HHMMSS>/... (and quality/<YYYYMMDD-...>).
    We pull the first 8-digit YYYYMMDD run-folder segment and format it as a date so
    it can be compared against the --date filter. Returns None if not found.
    """
    import re

    for seg in str(path).replace("\\", "/").split("/"):
        m = re.match(r"^(\d{4})(\d{2})(\d{2})-\d{6}", seg)
        if m:
            return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    return None


def _row_dates(row: dict[str, Any]) -> set[str]:
    """Return valid dates associated with a row (payload date and run-folder date)."""
    dates: set[str] = set()
    ts = str(row.get("started_at") or row.get("timestamp") or "")
    payload_date = ts[:10]
    try:
        datetime.strptime(payload_date, "%Y-%m-%d")
        dates.add(payload_date)
    except ValueError:
        pass
    folder_date = _folder_date_from_path(row.get("_source_path", ""))
    if folder_date:
        dates.add(folder_date)
    return dates


def filter_by_date_range(
    results: list[dict[str, Any]],
    from_date: str | None,
    to_date: str | None,
) -> list[dict[str, Any]]:
    """Keep rows with any associated date inside the inclusive range."""
    if not from_date and not to_date:
        return results

    def in_range(date_value: str) -> bool:
        return (not from_date or date_value >= from_date) and (not to_date or date_value <= to_date)

    return [row for row in results if any(in_range(d) for d in _row_dates(row))]


def filter_by_date(results: list[dict[str, Any]], date_str: str | None) -> list[dict[str, Any]]:
    """Backward-compatible exact-date filter."""
    return filter_by_date_range(results, date_str, date_str)


def get_all_headers(results: list[dict[str, Any]]) -> list[str]:
    """Extract union of all keys across all result dicts to prevent missing field errors.

    Internal bookkeeping keys (prefixed with '_', e.g. _source_path, _clean_*) are
    excluded so they don't leak into the exported CSV/Excel detail columns.
    """
    headers = []
    seen = set()
    for r in results:
        for k in r.keys():
            if k.startswith("_"):
                continue
            if k not in seen:
                seen.add(k)
                headers.append(k)
    return headers


def export_csv(results: list[dict[str, Any]], csv_path: Path) -> None:
    if not results:
        return
    headers = get_all_headers(results)
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for r in results:
            writer.writerow(r)
    print(f"[SUCCESS] Benchmark comparison CSV exported to: {csv_path.resolve()}")


def _v(r: dict[str, Any], *keys: str, default: float = 0.0) -> float:
    """Safely extract float metric from dict, handling None values and key fallbacks."""
    for key in keys:
        val = r.get(key)
        if val is not None:
            try:
                return float(val)
            except (TypeError, ValueError):
                continue
    return default


# EC2 on-demand $/hr price book (mirrors inference/load-test.py EC2_HOURLY_USD).
# Used ONLY by --fix-instance-labels to recompute cost when a row's recorded
# instance_type disagrees with the S3 run-folder hardware family.
_EC2_HOURLY_USD = {
    "g5.xlarge": 1.0060, "g5.2xlarge": 1.2120, "g5.4xlarge": 1.6240, "g5.8xlarge": 2.4480,
    "g5.12xlarge": 5.6720, "g5.16xlarge": 4.0960, "g5.24xlarge": 8.1440, "g5.48xlarge": 16.2880,
    "g6.xlarge": 0.8048, "g6.2xlarge": 0.9776, "g6.4xlarge": 1.3232, "g6.8xlarge": 2.0144,
    "g6.12xlarge": 4.6016, "g6.16xlarge": 3.3968, "g6.24xlarge": 6.6752, "g6.48xlarge": 13.3504,
    "g6e.xlarge": 1.8610, "g6e.2xlarge": 2.2420, "g6e.4xlarge": 3.0042, "g6e.8xlarge": 4.5286,
    "g6e.12xlarge": 10.6032, "g6e.16xlarge": 7.5772, "g6e.24xlarge": 15.0656, "g6e.48xlarge": 30.1312,
    "g7e.2xlarge": 3.3631, "g7e.4xlarge": 3.9982, "g7e.8xlarge": 5.2682,
    "g7e.12xlarge": 8.2861, "g7e.24xlarge": 16.5722, "g7e.48xlarge": 33.1443,
    "g7.2xlarge": 2.52, "trn1.32xlarge": 21.50, "trn2.48xlarge": 12.0,
}

# Known hardware families that can appear as an S3 run-folder segment.
_HW_FAMILIES = ("g5", "g6e", "g6", "g7e", "g7", "trn2", "trn1")


def _family_of_instance(instance_type: str) -> str | None:
    """'g6.12xlarge' -> 'g6', 'g7e.2xlarge' -> 'g7e'. None if empty."""
    it = str(instance_type or "").strip().lower()
    return it.split(".")[0] if it else None


def _folder_family_from_path(path: str) -> str | None:
    """The hardware family segment in an S3 path, e.g. results/<ts>/g6e/batch/... -> 'g6e'.

    Matches the exact folder segment against known families. Longer names first
    (g6e before g6) so 'g6e' is never mistaken for 'g6'.
    """
    segs = str(path).replace("\\", "/").split("/")
    for seg in segs:
        s = seg.strip().lower()
        for fam in _HW_FAMILIES:  # ordered longest-distinct first
            if s == fam:
                return fam
    return None


def fix_instance_labels(results: list[dict[str, Any]]) -> int:
    """One-time correction for rows whose recorded instance_type disagrees with the
    S3 run-folder hardware family (a pre-dynamic-instance mislabel).

    Rule:
      - Derive the TRUE family from the row's _source_path (results/<ts>/<family>/...).
      - Derive the LABELED family from row['instance_type'].
      - If they MATCH  -> leave the row exactly as-is (data is correct).
      - If they DIFFER -> the run really ran on <true family>; rewrite instance_type
        to the same SIZE under the true family (g6.12xlarge -> g6e.12xlarge) and
        RESCALE every price-derived cost field by (new_price / old_price), because
        all costs scale linearly with instance_hourly_usd. Performance metrics
        (throughput/latency/tokens) and energy_per_1m (power-based) are untouched.

    Returns the number of rows corrected.
    """
    _COST_FIELDS = (
        "cost_per_1m_tokens",
        "cost_per_1m_output_tokens",
        "cost_per_qa_form",
    )
    fixed = 0
    for r in results:
        true_fam = _folder_family_from_path(r.get("_source_path", ""))
        labeled = str(r.get("instance_type", "")).strip().lower()
        labeled_fam = _family_of_instance(labeled)
        if not true_fam or not labeled_fam or true_fam == labeled_fam:
            continue  # unknown folder, or already correct -> keep exact data

        # Same size under the true family: g6.12xlarge -> g6e.12xlarge
        size = labeled.split(".", 1)[1] if "." in labeled else ""
        corrected_instance = f"{true_fam}.{size}" if size else true_fam

        old_price = _EC2_HOURLY_USD.get(labeled)
        new_price = _EC2_HOURLY_USD.get(corrected_instance)
        if not new_price:
            # Can't price the corrected instance; still fix the label but warn.
            print(f"[FIX] {labeled} -> {corrected_instance} (label only; no price for {corrected_instance})")
            r["instance_type"] = corrected_instance
            fixed += 1
            continue

        r["instance_type"] = corrected_instance
        if old_price and old_price > 0:
            scale = new_price / old_price
            # Rescale price-derived cost fields WITHOUT re-rounding — keep full float
            # precision so tiny values (e.g. cost_per_qa_form ~0.001) are never lost.
            # Any rounding is done later at display time, not in the data.
            for fld in _COST_FIELDS:
                v = r.get(fld)
                if isinstance(v, (int, float)):
                    r[fld] = v * scale
            tpu = r.get("tokens_per_usd")
            if isinstance(tpu, (int, float)):
                r["tokens_per_usd"] = tpu / scale  # inverse: fewer tokens per $ at higher price
        r["instance_hourly_usd"] = new_price
        # Quality rows cache a display field derived from instance_type — refresh it.
        if "_clean_hw" in r:
            r["_clean_hw"] = corrected_instance
        fixed += 1
        print(f"[FIX] {labeled} -> {corrected_instance}  (cost x {new_price/old_price:.3f})" if old_price else
              f"[FIX] {labeled} -> {corrected_instance}")
    return fixed


def export_excel(
    perf_results: list[dict[str, Any]],
    qual_results: list[dict[str, Any]],
    xlsx_path: Path,
    sample_results: list[dict[str, Any]] | None = None,
    max_sample_rows: int = 2000,
) -> None:
    """Generates a formatted Excel (.xlsx) report with Performance, Quality,
    Quality Samples, and Detail sheets."""
    sample_results = sample_results or []
    if not perf_results and not qual_results and not sample_results:
        print("No benchmark or quality results to export.")
        return

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        import subprocess
        print("Installing openpyxl in Python environment...")
        subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"], check=False)
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            print("Notice: openpyxl unavailable — exporting CSV report.")
            export_csv(perf_results, xlsx_path.with_suffix(".csv"))
            return

    wb = openpyxl.Workbook()

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill_perf = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_fill_qual = PatternFill(start_color="276A3C", end_color="276A3C", fill_type="solid")

    # ── Sheet 1: Performance Comparison ────────────────────────────────────
    ws_perf = wb.active
    ws_perf.title = "Performance Benchmarks"

    headers_perf = [
        "Model (HF ID)",
        "Profile",
        "Concurrency",
        "TTFT p50 (ms)",
        "TTFT p95 (ms)",
        "ITL p95 (ms)",
        "Throughput (tok/s)",
        "Cost / 1M Tokens ($)",
        "Cost / Form ($)",
        "GPU Util (%)",
        "GPU Mem Used (MiB)",
        "KV Cache Usage (%)",
        "Instance Type",
        "Status",
    ]

    ws_perf.append(headers_perf)
    for col_num in range(1, len(headers_perf) + 1):
        cell = ws_perf.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill_perf
        cell.alignment = Alignment(horizontal="center", vertical="center")

    seen_perf = set()
    unique_perf = []
    for r in perf_results:
        key = (r.get("hf_id", ""), r.get("profile", ""), r.get("concurrency", 0), r.get("started_at", ""))
        if key not in seen_perf:
            seen_perf.add(key)
            unique_perf.append(r)

    for r in sorted(unique_perf, key=lambda x: (x.get("hf_id", ""), x.get("profile", ""), x.get("concurrency", 0))):
        ws_perf.append([
            r.get("hf_id", "unknown").split("/")[-1],
            r.get("profile", "unknown"),
            r.get("concurrency", 1),
            round(_v(r, "ttft_p50_ms", "median_ttft_ms"), 1),
            round(_v(r, "ttft_p95_ms", "p95_ttft_ms"), 1),
            round(_v(r, "itl_p95_ms", "p95_itl_ms"), 1),
            round(_v(r, "throughput_tokens_s", "output_throughput_tokens_s"), 1),
            round(_v(r, "cost_per_1m_tokens"), 4),
            round(_v(r, "cost_per_qa_form"), 6),
            round(_v(r, "gpu_utilization_pct"), 1),
            round(_v(r, "gpu_memory_used_mib", "gpu_mem_used_mib"), 0),
            round(_v(r, "gpu_cache_usage_pct", "kv_cache_utilization_pct"), 1),
            r.get("instance_type", "unknown"),
            "PASSED" if r.get("status") == "passed" else "SLO FAIL",
        ])

    # ── Sheet 2: Quality & Accuracy Evaluation ──────────────────────────────
    ws_qual = wb.create_sheet(title="Quality Evaluation")
    # Confusion cells, unparseable and truncated counts come straight from the
    # evaluator's result row (quality-eval.py writes them); nothing is derived here.
    headers_qual = [
        "Model",
        "Instance Type / HW",
        "Quantization",
        "Total Transcripts",
        "Accuracy (%)",
        "Precision",
        "Recall",
        "AutoQA F1 Score",
        "Macro F1",
        "TP",
        "FP",
        "FN",
        "TN",
        "Unparseable (gold Yes)",
        "Unparseable (gold No)",
        "Other Unscored",
        "Matrix Total",
        "Unparseable",
        "Truncated",
        "Recovered by Retry",
        "Verdict Sources",
        "Decoding Mode",
        "Timestamp",
    ]

    ws_qual.append(headers_qual)
    for col_num in range(1, len(headers_qual) + 1):
        cell = ws_qual.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill_qual
        cell.alignment = Alignment(horizontal="center", vertical="center")

    seen_qual = set()
    unique_qual = []
    for q in qual_results:
        key = (q.get("_clean_model", ""), q.get("_clean_quant", ""), q.get("_clean_hw", ""), q.get("started_at", ""))
        if key not in seen_qual:
            seen_qual.add(key)
            unique_qual.append(q)

    for q in sorted(unique_qual, key=lambda x: (x.get("_clean_model", ""), x.get("_clean_quant", ""))):
        conf = q.get("confusion") or {}
        if not isinstance(conf, dict):
            conf = {}
        ws_qual.append([
            q.get("_clean_model", "gpt-oss-20b"),
            q.get("_clean_hw", "-"),
            q.get("_clean_quant", "-"),
            q.get("_clean_rows", 1200),
            round(q.get("_clean_acc", 0.0), 2),
            round(_v(q, "precision"), 4),
            round(_v(q, "recall"), 4),
            round(q.get("_clean_f1", 0.0), 4),
            # Absent field shows "-" rather than 0, so an older result row is not
            # read as a genuine macro-F1 of zero.
            round(_v(q, "macro_f1"), 4) if q.get("macro_f1") is not None else "-",
            conf.get("tp", "-"),
            conf.get("fp", "-"),
            conf.get("fn", "-"),
            conf.get("tn", "-"),
            conf.get("unparseable_pos", "-"),
            conf.get("unparseable_neg", "-"),
            conf.get("unscored_other", "-"),
            q.get("confusion_total", "-"),
            q.get("n_unparseable", "-"),
            q.get("n_truncated", "-"),
            q.get("n_fallback", "-"),
            # Compact provenance string, e.g. "answer=120 conclusion=940 retry=140".
            (
                " ".join(f"{k}={v}" for k, v in sorted((q.get("verdict_sources") or {}).items()))
                if isinstance(q.get("verdict_sources"), dict) and q.get("verdict_sources")
                else "-"
            ),
            q.get("decoding_mode", "-"),
            str(q.get("started_at", q.get("timestamp", "-")))[:19].replace("T", " "),
        ])

    # ── Sheet 3: Quality Samples (per-row model output) ─────────────────────
    # Present only when a run was submitted with --dump-samples. Every column is
    # a field written by quality-eval.py; the report does not re-judge anything.
    ws_samples = None
    if sample_results:
        ws_samples = wb.create_sheet(title="Quality Samples")
        headers_samples = [
            "Model",
            "Instance Type / HW",
            "Quantization",
            "Run",
            "Data ID",
            "Question (rubric)",
            "Prompt Chars",
            "Expected (gold)",
            "Actual (predicted)",
            "Match",
            "Truncated",
            "Finish Reason",
            "Retry Used",
            "Verdict From",
            "Model Answer",
            "Model Reasoning",
        ]
        # The input prompt is only present when a run used --dump-inputs, so the
        # column appears only when at least one row actually carries it.
        has_inputs = any(s.get("input_prompt") for s in sample_results)
        if has_inputs:
            headers_samples.append("Input Prompt (transcript)")
        ws_samples.append(headers_samples)
        for col_num in range(1, len(headers_samples) + 1):
            cell = ws_samples.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = PatternFill(start_color="7F4F24", end_color="7F4F24", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Excel caps a cell at 32767 characters. Reasoning is trimmed to stay
        # readable; the input prompt gets a much larger budget because a partial
        # transcript cannot be used to audit a verdict.
        cell_cap = 4000
        input_cap = 32000

        def _fit(value: Any, cap: int) -> str:
            text = str(value or "")
            return text if len(text) <= cap else text[:cap] + " ...[truncated for Excel]"

        written = 0
        for s in sorted(
            sample_results,
            key=lambda x: (x.get("_clean_model", ""), x.get("_run", ""), str(x.get("data_id", ""))),
        ):
            if written >= max_sample_rows:
                break
            # Newer runs split the reply into answer + reasoning. Older sample
            # files only have the glued "output", which is shown as the answer.
            answer = _fit(s.get("answer", s.get("output", "")), cell_cap)
            reasoning = _fit(s.get("reasoning", ""), cell_cap)
            match_val = s.get("match")
            row_cells = [
                s.get("_clean_model", "-"),
                s.get("_clean_hw", "-"),
                s.get("_clean_quant", "-"),
                s.get("_run", "-"),
                s.get("data_id", "-"),
                str(s.get("question", "") or "-"),
                s.get("prompt_chars", "-") or "-",
                s.get("gold", "-"),
                s.get("pred", "-"),
                ("MATCH" if match_val else "MISMATCH") if match_val is not None else "-",
                # "-" when the sample predates the truncation field, so an unknown
                # state is never displayed as confirmed not-truncated.
                ("YES" if s.get("truncated") else "no") if s.get("truncated") is not None else "-",
                s.get("finish_reason", "-"),
                ("YES" if s.get("fallback") else "no") if s.get("fallback") is not None else "-",
                s.get("verdict_source", "-") or "-",
                answer or "-",
                reasoning or "-",
            ]
            if has_inputs:
                row_cells.append(_fit(s.get("input_prompt", ""), input_cap) or "-")
            ws_samples.append(row_cells)
            written += 1

        if len(sample_results) > written:
            print(
                f"[NOTE] Quality Samples sheet capped at {written} of "
                f"{len(sample_results)} rows (raise --max-sample-rows to include more)."
            )

    # ── Sheet 4: All Detailed Runs ──────────────────────────────────────────
    ws_detail = wb.create_sheet(title="All Detailed Perf Runs")
    if unique_perf:
        all_keys = get_all_headers(unique_perf)
        ws_detail.append(all_keys)

        for col_num in range(1, len(all_keys) + 1):
            cell = ws_detail.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

        for r in unique_perf:
            row_vals = []
            for k in all_keys:
                v = r.get(k, "")
                if isinstance(v, (list, dict)):
                    row_vals.append(json.dumps(v))
                else:
                    row_vals.append(v)
            ws_detail.append(row_vals)

    # Auto-adjust column widths. Capped so a long model-output cell cannot stretch
    # a column past the width of the screen.
    for ws in [s for s in (ws_perf, ws_qual, ws_samples, ws_detail) if s is not None]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 80)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(f"\n[SUCCESS] Model comparison Excel report generated successfully!")
    print(f"[REPORT]  File Location: {xlsx_path.resolve()}\n")


def print_comparison_table(
    perf_results: list[dict[str, Any]],
    qual_results: list[dict[str, Any]],
    sample_results: list[dict[str, Any]] | None = None,
) -> None:
    sample_results = sample_results or []
    sep = "=" * 105
    print(f"\n{sep}")
    print(
        f"  LLM BENCHMARK MODEL COMPARISON SUMMARY ({len(perf_results)} perf runs, "
        f"{len(qual_results)} quality runs, {len(sample_results)} quality sample rows)"
    )
    print(sep)

    # ── Section 1: Load Test Performance Results ─────────────────────────────
    if perf_results:
        print("\n  [ PERFORMANCE & EFFICIENCY BENCHMARKS ]")
        print(
            f"  {'Model':<22} {'HW':<8} {'Profile':<9} {'Concur':>6} "
            f"{'TTFT p95':>10} {'ITL p95':>9} {'Tok/s':>9} "
            f"{'Cost/1M':>9} {'GPU Util':>9} {'Status':<9}"
        )
        print(f"  {'-'*101}")

        seen = set()
        unique_perf = []
        for r in perf_results:
            key = (r.get("hf_id", ""), r.get("profile", ""), r.get("concurrency", 0), r.get("started_at", ""))
            if key not in seen:
                seen.add(key)
                unique_perf.append(r)

        for r in sorted(unique_perf, key=lambda x: (x.get("hf_id", ""), x.get("profile", ""), x.get("concurrency", 0))):
            model = r.get("hf_id", "unknown").split("/")[-1]
            hw = r.get("instance_type", r.get("hardware", "-"))
            profile = r.get("profile", "unknown")
            concurrency = r.get("concurrency", 1)
            ttft_p95 = f"{_v(r, 'ttft_p95_ms', 'p95_ttft_ms'):.1f}ms"
            itl_p95 = f"{_v(r, 'itl_p95_ms', 'p95_itl_ms'):.1f}ms"
            throughput = f"{_v(r, 'throughput_tokens_s', 'output_throughput_tokens_s'):.1f}"
            cost = f"${_v(r, 'cost_per_1m_tokens'):.4f}"
            gpu_util = f"{_v(r, 'gpu_utilization_pct'):.1f}%"
            status = "PASSED" if r.get("status") == "passed" else "SLO FAIL"

            print(
                f"  {model:<22} {hw:<8} {profile:<9} {concurrency:>6} "
                f"{ttft_p95:>10} {itl_p95:>9} {throughput:>9} "
                f"{cost:>9} {gpu_util:>9} {status:<9}"
            )

    # ── Section 2: Quality Evaluation Results ────────────────────────────────
    if qual_results:
        print("\n  [ QUALITY & ACCURACY BENCHMARKS (AutoQA F1 Score) ]")
        print(
            f"  {'Model':<22} {'HW':<8} {'Quant':<8} {'Rows':>6} "
            f"{'Accuracy':>10} {'Precision':>10} {'Recall':>9} {'F1 Score':>10} "
            f"{'Unparse':>8} {'Trunc':>6}"
        )
        print(f"  {'-'*106}")

        seen_q = set()
        unique_qual = []
        for q in qual_results:
            key = (q.get("_clean_model", ""), q.get("_clean_quant", ""), q.get("_clean_hw", ""), q.get("started_at", ""))
            if key not in seen_q:
                seen_q.add(key)
                unique_qual.append(q)

        for q in sorted(unique_qual, key=lambda x: (x.get("_clean_model", ""), x.get("_clean_quant", ""))):
            model = str(q.get("_clean_model", "gpt-oss-20b"))
            hw = str(q.get("_clean_hw", "-"))
            quant = str(q.get("_clean_quant", "-"))
            rows = q.get("_clean_rows", 1200)
            acc = f"{q.get('_clean_acc', 0.0):.2f}%"
            prec = f"{_v(q, 'precision'):.4f}"
            rec = f"{_v(q, 'recall'):.4f}"
            f1 = f"{q.get('_clean_f1', 0.0):.4f}"
            unparse = str(q.get("n_unparseable", "-"))
            trunc = str(q.get("n_truncated", "-"))

            print(
                f"  {model:<22} {hw:<8} {quant:<8} {rows:>6} "
                f"{acc:>10} {prec:>10} {rec:>9} {f1:>10} "
                f"{unparse:>8} {trunc:>6}"
            )

            # Confusion matrix straight from the evaluator, with the reconciliation
            # check so a matrix that does not add up to the row count is obvious.
            conf = q.get("confusion") if isinstance(q.get("confusion"), dict) else None
            if conf:
                total = q.get("confusion_total")
                recon = ""
                if isinstance(total, int):
                    recon = f"   total={total}/{rows} {'OK' if total == rows else 'MISMATCH'}"
                print(
                    f"      confusion: TP={conf.get('tp', '-')} FP={conf.get('fp', '-')} "
                    f"FN={conf.get('fn', '-')} TN={conf.get('tn', '-')} "
                    f"unparse(Yes)={conf.get('unparseable_pos', '-')} "
                    f"unparse(No)={conf.get('unparseable_neg', '-')}{recon}"
                )
            if isinstance(q.get("n_truncated"), int) and q["n_truncated"] > 0:
                print(
                    f"      NOTE: {q['n_truncated']} repl{'y' if q['n_truncated'] == 1 else 'ies'} "
                    f"hit the token budget — treat this run as an integration issue, "
                    f"not model quality."
                )

    print(f"\n{sep}\n")


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare LLM benchmark result runs")
    parser.add_argument(
        "--results-dir",
        default="results",
        help="Path to directory containing benchmark result files (default: results/)",
    )
    parser.add_argument(
        "--export-excel",
        default="results/model_comparison.xlsx",
        help="Path to save output Excel report (default: results/model_comparison.xlsx)",
    )
    parser.add_argument(
        "--export-csv",
        default="",
        help="Optional path to export comparison CSV (e.g. results/model_comparison.csv)",
    )
    parser.add_argument(
        "--s3-bucket",
        default="",
        help="Optional S3 bucket to sync historical results from (e.g. shellkode-ai-results)",
    )
    parser.add_argument(
        "--date",
        default=None,
        help=(
            "Include one exact date based on started_at or S3 run-folder timestamp. "
            "Use 'today' (UTC) or YYYY-MM-DD. Cannot be combined with date-range options."
        ),
    )
    parser.add_argument(
        "--from-date",
        default=None,
        help=(
            "Inclusive range start: 'today' or YYYY-MM-DD. If --to-date is omitted, "
            "the range ends today (UTC)."
        ),
    )
    parser.add_argument(
        "--to-date",
        default=None,
        help="Inclusive range end: 'today' or YYYY-MM-DD.",
    )
    parser.add_argument(
        "--max-sample-rows",
        type=int,
        default=2000,
        help=(
            "Cap on rows written to the Quality Samples sheet (default 2000). "
            "Samples exist only for runs launched with --dump-samples."
        ),
    )
    parser.add_argument(
        "--fix-instance-labels",
        action="store_true",
        help=(
            "One-time correction: for any row whose recorded instance_type family "
            "disagrees with its S3 run-folder family (results/<ts>/<family>/...), "
            "rewrite instance_type to the folder's family (same size) and recompute "
            "cost from the correct instance price. Rows that already match are left "
            "exactly as-is. Performance metrics are never changed."
        ),
    )
    args = parser.parse_args()

    results_path = Path(args.results_dir)

    if args.s3_bucket:
        sync_results_from_s3(args.s3_bucket, results_path)

    perf_results = load_results_local(results_path)
    qual_results = load_quality_results(results_path)
    sample_results = load_quality_samples(results_path)

    if args.fix_instance_labels:
        n = fix_instance_labels(perf_results) + fix_instance_labels(qual_results)
        print(f"[FIX] Corrected {n} mislabeled row(s) (instance_type + cost) from S3 folder family.")

    from_date, to_date = resolve_date_range(args.date, args.from_date, args.to_date)
    if from_date or to_date:
        before_perf, before_qual = len(perf_results), len(qual_results)
        before_samples = len(sample_results)
        perf_results = filter_by_date_range(perf_results, from_date, to_date)
        qual_results = filter_by_date_range(qual_results, from_date, to_date)
        # Sample rows have no started_at of their own, so they are matched on the
        # S3 run-folder date by the same range filter.
        sample_results = filter_by_date_range(sample_results, from_date, to_date)
        label = from_date if from_date == to_date else f"{from_date or 'earliest'} through {to_date or 'latest'}"
        print(
            f"[FILTER] Date = {label} (inclusive): "
            f"perf runs {before_perf} -> {len(perf_results)}, "
            f"quality runs {before_qual} -> {len(qual_results)}, "
            f"quality samples {before_samples} -> {len(sample_results)}"
        )

    print_comparison_table(perf_results, qual_results, sample_results)

    if args.export_csv and perf_results:
        export_csv(perf_results, Path(args.export_csv))

    if args.export_excel:
        export_excel(
            perf_results,
            qual_results,
            Path(args.export_excel),
            sample_results=sample_results,
            max_sample_rows=args.max_sample_rows,
        )


if __name__ == "__main__":
    main()