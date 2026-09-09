#!/usr/bin/env python3
"""
scripts/compare_models.py

Parses all JSONL benchmark results from results/ directory (recursively)
or syncs from S3, then generates a multi-model comparison table and Excel/CSV report.

Usage:
  python scripts/compare_models.py
  python scripts/compare_models.py --s3-bucket shellkode-ai-results
  python scripts/compare_models.py --export-excel results/model_comparison.xlsx
"""

import argparse
import csv
import json
import os
import sys
from pathlib import Path
from typing import Any


def load_results_local(results_dir: Path) -> list[dict[str, Any]]:
    rows = []
    if not results_dir.exists():
        return rows

    # Use rglob to search recursively through all subdirectories
    for f in results_dir.rglob("*.jsonl"):
        if f.name.startswith("comparison") or f.name.startswith("summary"):
            continue
        with f.open("r", encoding="utf-8") as file:
            for line in file:
                if line.strip():
                    try:
                        rows.append(json.loads(line))
                    except json.JSONDecodeError:
                        pass
    return rows


def sync_results_from_s3(bucket: str, local_dir: Path) -> None:
    """Sync all benchmark results from S3 bucket into local results/ directory."""
    import subprocess

    cmd = f"aws s3 sync s3://{bucket}/results/ {local_dir}/"
    print(f"Syncing benchmark results from S3 (s3://{bucket}/results/) ...")
    try:
        subprocess.run(cmd, shell=True, check=True)
        print("S3 sync complete.")
    except Exception as e:
        print(f"Warning: S3 sync failed or AWS CLI not found: {e}")


def get_all_headers(results: list[dict[str, Any]]) -> list[str]:
    """Extract union of all keys across all result dicts to prevent missing field errors."""
    headers = []
    seen = set()
    for r in results:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                headers.append(k)
    return headers


def export_csv(results: list[dict[str, Any]], csv_path: Path) -> None:
    if not results:
        return
    headers = get_all_headers(results)
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for r in results:
            writer.writerow(r)
    print(f"[SUCCESS] Benchmark comparison CSV exported to: {csv_path.resolve()}")


def _v(r: dict[str, Any], *keys: str, default: float = 0.0) -> float:
    """Safely extract float metric from dict, handling None values and key fallbacks."""
    for key in keys:
        val = r.get(key)
        if val is not None:
            try:
                return float(val)
            except (TypeError, ValueError):
                continue
    return default


def export_excel(results: list[dict[str, Any]], xlsx_path: Path) -> None:
    """Generates a formatted Excel (.xlsx) report with a Summary sheet and Details sheet."""
    if not results:
        print("No benchmark results to export.")
        return

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        import subprocess
        print("Installing openpyxl in Python environment...")
        subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"], check=False)
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            print("Notice: openpyxl unavailable — exporting CSV report.")
            export_csv(results, xlsx_path.with_suffix(".csv"))
            return

    wb = openpyxl.Workbook()

    # ── Sheet 1: Comparison Summary ─────────────────────────────────────────
    ws_summary = wb.active
    ws_summary.title = "Model Comparison"

    headers_summary = [
        "Model (HF ID)",
        "Profile",
        "Concurrency",
        "TTFT p50 (ms)",
        "TTFT p95 (ms)",
        "ITL p95 (ms)",
        "Throughput (tok/s)",
        "Cost / 1M Tokens ($)",
        "Cost / Form ($)",
        "GPU Util (%)",
        "GPU Mem Used (MiB)",
        "KV Cache Usage (%)",
        "Instance Type",
        "Status",
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

    ws_summary.append(headers_summary)
    for col_num in range(1, len(headers_summary) + 1):
        cell = ws_summary.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Deduplicate rows by (hf_id, profile, concurrency, started_at)
    seen = set()
    unique_results = []
    for r in results:
        key = (r.get("hf_id", ""), r.get("profile", ""), r.get("concurrency", 0), r.get("started_at", ""))
        if key not in seen:
            seen.add(key)
            unique_results.append(r)

    for r in sorted(unique_results, key=lambda x: (x.get("hf_id", ""), x.get("profile", ""), x.get("concurrency", 0))):
        ws_summary.append([
            r.get("hf_id", "unknown").split("/")[-1],
            r.get("profile", "unknown"),
            r.get("concurrency", 1),
            round(_v(r, "ttft_p50_ms", "median_ttft_ms"), 1),
            round(_v(r, "ttft_p95_ms", "p95_ttft_ms"), 1),
            round(_v(r, "itl_p95_ms", "p95_itl_ms"), 1),
            round(_v(r, "throughput_tokens_s", "output_throughput_tokens_s"), 1),
            round(_v(r, "cost_per_1m_tokens"), 4),
            round(_v(r, "cost_per_qa_form"), 6),
            round(_v(r, "gpu_utilization_pct"), 1),
            round(_v(r, "gpu_memory_used_mib", "gpu_mem_used_mib"), 0),
            round(_v(r, "gpu_cache_usage_pct", "kv_cache_utilization_pct"), 1),
            r.get("instance_type", "unknown"),
            "PASSED" if r.get("status") == "passed" else "SLO FAIL",
        ])

    # ── Sheet 2: All Detailed Runs ──────────────────────────────────────────
    ws_detail = wb.create_sheet(title="All Detailed Runs")
    if unique_results:
        all_keys = get_all_headers(unique_results)
        ws_detail.append(all_keys)

        for col_num in range(1, len(all_keys) + 1):
            cell = ws_detail.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

        for r in unique_results:
            row_vals = []
            for k in all_keys:
                v = r.get(k, "")
                if isinstance(v, (list, dict)):
                    row_vals.append(json.dumps(v))
                else:
                    row_vals.append(v)
            ws_detail.append(row_vals)

    # Auto-adjust column widths
    for ws in [ws_summary, ws_detail]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(f"\n[SUCCESS] Model comparison Excel report generated successfully!")
    print(f"[REPORT]  File Location: {xlsx_path.resolve()}\n")


def print_comparison_table(results: list[dict[str, Any]]) -> None:
    if not results:
        print("No benchmark result JSONL files found.")
        return

    # Deduplicate rows by (hf_id, profile, concurrency, started_at)
    seen = set()
    unique_results = []
    for r in results:
        key = (r.get("hf_id", ""), r.get("profile", ""), r.get("concurrency", 0), r.get("started_at", ""))
        if key not in seen:
            seen.add(key)
            unique_results.append(r)

    sep = "=" * 100
    print(f"\n{sep}")
    print(f"  LLM BENCHMARK MODEL COMPARISON SUMMARY ({len(unique_results)} runs found)")
    print(sep)
    print(
        f"  {'Model (HF ID)':<26} {'Profile':<10} {'Concur':>6} "
        f"{'TTFT p95':>10} {'ITL p95':>9} {'Tok/s':>9} "
        f"{'Cost/1M':>10} {'GPU Util':>9} {'Status':<10}"
    )
    print(f"  {'-'*96}")

    for r in sorted(unique_results, key=lambda x: (x.get("hf_id", ""), x.get("profile", ""), x.get("concurrency", 0))):
        hf_id = r.get("hf_id", "unknown").split("/")[-1]
        profile = r.get("profile", "unknown")
        concurrency = r.get("concurrency", 1)
        ttft_p95 = f"{_v(r, 'ttft_p95_ms', 'p95_ttft_ms'):.1f}ms"
        itl_p95 = f"{_v(r, 'itl_p95_ms', 'p95_itl_ms'):.1f}ms"
        throughput = f"{_v(r, 'throughput_tokens_s', 'output_throughput_tokens_s'):.1f}"
        cost = f"${_v(r, 'cost_per_1m_tokens'):.4f}"
        gpu_util = f"{_v(r, 'gpu_utilization_pct'):.1f}%"
        status = "PASSED" if r.get("status") == "passed" else "SLO FAIL"

        print(
            f"  {hf_id:<26} {profile:<10} {concurrency:>6} "
            f"{ttft_p95:>10} {itl_p95:>9} {throughput:>9} "
            f"{cost:>10} {gpu_util:>9} {status:<10}"
        )

    print(sep)


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare LLM benchmark result runs")
    parser.add_argument(
        "--results-dir",
        default="results",
        help="Path to directory containing benchmark result files (default: results/)",
    )
    parser.add_argument(
        "--export-excel",
        default="results/model_comparison.xlsx",
        help="Path to save output Excel report (default: results/model_comparison.xlsx)",
    )
    parser.add_argument(
        "--export-csv",
        default="",
        help="Optional path to export comparison CSV (e.g. results/model_comparison.csv)",
    )
    parser.add_argument(
        "--s3-bucket",
        default="",
        help="Optional S3 bucket to sync historical results from (e.g. shellkode-ai-results)",
    )
    args = parser.parse_args()

    results_path = Path(args.results_dir)

    if args.s3_bucket:
        sync_results_from_s3(args.s3_bucket, results_path)

    results = load_results_local(results_path)
    print_comparison_table(results)

    if args.export_csv and results:
        export_csv(results, Path(args.export_csv))

    if args.export_excel and results:
        export_excel(results, Path(args.export_excel))


if __name__ == "__main__":
    main()
